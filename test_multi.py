"""
API-free tests for the triage and multi-document pipeline.

    python test_multi.py

Everything here is FICTIONAL test data (not any real trust). No API key is
needed: the two LLM calls are replaced with canned annotations when main.py's
routing is tested.
"""

import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("ANTHROPIC_API_KEY", "test-key-not-used")

from openpyxl import load_workbook

import cli as cli_module
import service
from cell_map import FIXED_ROWS
from settings import load_settings

SETTINGS = load_settings()
ANNOTATION_TEMPLATE_FILE = str(SETTINGS.template_single)
AMENDMENT_TEMPLATE_FILE = str(SETTINGS.template_multi)
from packet import MAX_AMENDMENTS, build_packet, triage
from render_multi import render_packet_to_workbook
from schema import (
    BequestLanguage, DefinedTerm, DocumentAnnotation, Finding as SingleFinding,
    IncomeBeneficiaries, RestrictionsContingencies, UMMention,
)
from schema_multi import (
    AmendmentEntry, BequestLanguageHistory, DefinedTermHistory, Finding, FieldHistory,
    IncomeBeneficiariesHistory, InstrumentInfo, PacketAnnotation,
    RestrictionsContingenciesHistory, reconcile_with_packet,
)
from schema import UncategorizedFinding

TMP = Path(tempfile.mkdtemp())
passed = 0


def check(condition, message):
    global passed
    assert condition, f"FAILED: {message}"
    passed += 1


def write(name, text="Some instrument text."):
    p = TMP / name
    p.write_text(text)
    return str(p)


# ---------------------------------------------------------------------------
# 1. Triage and packet validation
# ---------------------------------------------------------------------------
trust = write("trust.txt")
a1, a2, a4 = write("a1.txt"), write("a2.txt"), write("a4.txt")

p = build_packet("T", "Trust", trust, [])
check(triage(p) == "single", "no amendments -> single")

p = build_packet("T", "Trust", trust, [{"number": 1, "label": "First", "path": a1}])
check(triage(p) == "multi", "one amendment -> multi")

p = build_packet("T", "Trust", trust, [
    {"number": 4, "label": "Fourth", "path": a4},
    {"number": 1, "label": "First", "path": a1},
    {"number": 2, "label": "Second", "path": a2},
])
check(p.amendment_numbers == [1, 2, 4], "amendments are sorted by number, not list order")
check(p.missing_numbers == [3], "gap in the chain is detected (3 is missing)")


def raises(fn, fragment):
    try:
        fn()
    except ValueError as e:
        return fragment in str(e)
    return False


check(raises(lambda: build_packet("T", "Trust", trust, [
    {"number": 1, "label": "A", "path": a1}, {"number": 1, "label": "B", "path": a2}]),
    "more than once"), "duplicate amendment numbers are rejected")
check(raises(lambda: build_packet("T", "Trust", trust, [{"number": 0, "label": "A", "path": a1}]),
             ">= 1"), "amendment number 0 is rejected")
check(raises(lambda: build_packet("T", "Trust", trust, [{"number": 1, "label": "A", "path": "/nope.txt"}]),
             "not found"), "missing file is rejected")
check(raises(lambda: build_packet("T", "Trust", trust, [{"number": 1, "label": "A"}]),
             "missing"), "incomplete entry is rejected")
check(raises(lambda: build_packet("T", "Trust", trust, [
    {"number": i, "label": str(i), "path": a1} for i in range(1, MAX_AMENDMENTS + 2)]),
    "room for"), "more amendments than the template holds is rejected")
check(raises(lambda: build_packet("T", "Trust", write("empty.txt", "  \n"), []),
             "empty"), "empty trust file is rejected")

# ---------------------------------------------------------------------------
# 2. Fixtures: a fictional packet annotation
# ---------------------------------------------------------------------------

from testkit import hist, make_packet_annotation, tiny_single_annotation


# ---------------------------------------------------------------------------
# 3. reconcile_with_packet
# ---------------------------------------------------------------------------
ann = make_packet_annotation()
# scramble amendment 2/3 order, drop amendment 3's entry on one field, add a bogus entry
h = ann.bequest_language.um_receives
h.amendments = [h.amendments[1], h.amendments[0], AmendmentEntry(
    amendment_number=9, effect="Changed", description="bogus", citation="")]
warnings = reconcile_with_packet(ann, [1, 2, 3])
check([e.amendment_number for e in h.amendments] == [1, 2, 3], "entries put back in packet order")
check(h.amendments[2].effect == "Could not determine", "a skipped amendment entry becomes a visible 'Could not determine'")
check(any("unexpected" in w for w in warnings) and any("no entry for amendment 3" in w for w in warnings),
      "reconcile reports both the missing and the unexpected entry")

check(AmendmentEntry(amendment_number=1, effect="no CHANGE", description="", citation="").effect == "No change",
      "effect casing is normalized")
check(Finding(status="not in this Document").status == "Not in this document", "status casing is normalized")
check(InstrumentInfo(amendment_number=1, form="Changes_Only", revokes=[], note="").form == "changes_only",
      "form casing is normalized")

# ---------------------------------------------------------------------------
# 4. Rendering, with overflow in EVERY variable section
# ---------------------------------------------------------------------------
pk = build_packet("TEST — Fictional Trust with 3 Amendments", "Fictional Trust (2001)", trust, [
    {"number": 1, "label": "First Amendment (2003)", "path": a1},
    {"number": 2, "label": "Second Amendment (2004)", "path": a2},
    {"number": 3, "label": "Third Amendment (2006)", "path": write("a3.txt")},
])
ann = make_packet_annotation(n_terms=14, n_mentions=13, n_uncat=7, n_judgment=11, n_questions=8)
out = str(TMP / "rendered.xlsx")
render_packet_to_workbook(ann, pk, AMENDMENT_TEMPLATE_FILE, out)
ws = load_workbook(out)["Annotation"]

# fixed rows land where the template has them
r = FIXED_ROWS["bequest_language.um_receives"]
check(ws[f"B{r}"].value == "25% of the residue" and ws[f"D{r}"].value == "Found", "original block written")
check(ws[f"G{r}"].value == "No change", "amendment 1: 'No change'")
check(ws[f"H{r}"].value == "Changed: now 10% of the residue (Second Amendment, §3, p. 4)", "amendment 2: composed text")
check(ws[f"P{r}"].value == "10% of the residue" and ws[f"Q{r}"].value == "Second Amendment, §3, p. 4",
      "operative answer + controlling provision written")
r = FIXED_ROWS["bequest_language.trigger_event"]
check(ws[f"R{r}"].value == "Ambiguous — see note", "'Ambiguous' maps to the template's dropdown string")

# headers + hidden columns
check(ws["G11"].value == "First Amendment (2003)\nchanges only — REVOKED by Amendment 2",
      "amendment 1 header: form, and that amendment 2 revokes it")
check(ws["H11"].value == "Second Amendment (2004)\nfull restatement", "amendment 2 header")
check(ws["I11"].value == "Third Amendment (2006)\nchanges only", "amendment 3 header")

hidden = {k: d.hidden for k, d in ws.column_dimensions.items()}
check(not any(hidden.get(c) for c in "GHI"), "used amendment columns stay visible")
check(all(hidden.get(c) for c in "JKLMN"), "unused amendment columns (J-N) are hidden")

# section shifts: terms +2, mentions +1, uncategorized +1, judgment +1, questions +1
def find_row(prefix):
    for row in range(1, ws.max_row + 1):
        v = ws[f"A{row}"].value
        if isinstance(v, str) and v.startswith(prefix):
            return row
    raise AssertionError(f"row starting {prefix!r} not found")

band4, band5, band6, band7, band8, band9 = (find_row(p) for p in
    ("4 · DEFINED", "5 · ALL MENTIONS", "6 · FOUND", "7 · JUDGMENT", "8 · QUESTIONS", "9 · FEEDBACK"))
check(band4 == 28, "section 4 header has not moved")
check(band5 == 41 + 2, "section 5 header moved down by the 2 inserted term rows")
check(band6 == 54 + 3, "section 6 header moved by terms+mentions inserts (3)")
check(band7 == 61 + 4, "section 7 header moved by 4")
check(band8 == 72 + 5, "section 8 header moved by 5")
check(band9 == 80 + 6, "section 9 header moved by 6")

check(ws[f"A{band4 + 14}"].value == "Term 14" and ws[f"B{band4 + 14}"].value == "definition 14",
      "14th defined term written into an inserted row")
check(ws[f"A{band5 + 1}"].value == "Mention 1" and ws[f"B{band5 + 1}"].value == "mention text 1",
      "mentions start right under their (shifted) header, not on top of it")
check(ws[f"A{band5 + 13}"].value == "Mention 13", "13th mention written into an inserted row")
check(ws[f"B{band6 + 7}"].value == "uncategorized 7", "7th uncategorized finding written")
check(ws[f"B{band7 + 11}"].value == "judgment call 11", "11th judgment call written")
check(ws[f"B{band8 + 8}"].value == "question 8?", "8th question written")
check(ws[f"B{band9 + 1}"].value is None, "feedback section stays empty")

# machinery follows the rows
dv_ranges = {re.match(r"[A-Z]+", str(d.sqref)).group(0): str(d.sqref) for d in ws.data_validations.dataValidation}
check(f"D{band4 + 1}:D{band4 + 14}" in dv_ranges["D"], "status dropdown (original) covers the 14 term rows")
check(f"R{band5 + 1}:R{band5 + 13}" in dv_ranges["R"], "status dropdown (operative) covers the 13 mention rows")
check(ws["G9"].value == f'=COUNTIF(G13:G{band5 + 13},"*")-COUNTIF(G13:G{band5 + 13},"No change*")',
      "amendment tally formula extended to the last mention row")
check(ws.row_dimensions[band5].height == 19.5, "section header row kept its height after the shift")

# ---------------------------------------------------------------------------
# 5. Formulas calculate (recalc.py) and the amendment tally counts correctly
# ---------------------------------------------------------------------------
res = subprocess.run(
    [sys.executable, "/mnt/skills/public/xlsx/scripts/recalc.py", out],
    capture_output=True, text=True)
if res.returncode == 0 and '"status": "success"' in res.stdout:
    v = load_workbook(out, data_only=True)["Annotation"]
    check(v["G9"].value == 1, "amendment 1 tally: 1 field changed/flagged (Added: tax-exempt)")
    check(v["H9"].value == 1, "amendment 2 tally: 1 (Changed: 10%)")
    check(v["I9"].value == 1, "amendment 3 tally: 1 (Ambiguous: trigger event)")
    check(v["S9"].value == 1, "operative tally counts the one 'Ambiguous — see note'")
else:
    print("(recalc.py unavailable here; skipped formula-value checks)")

# ---------------------------------------------------------------------------
# 6. main.py routing, with the two LLM calls stubbed out
# ---------------------------------------------------------------------------

calls = []
service.analyze_document = lambda *a, **k: calls.append("single") or tiny_single_annotation()
service.analyze_packet = lambda packet, **k: calls.append("multi") or make_packet_annotation(
    n_terms=1, n_mentions=1)

cli_module.OUTPUT_DIR = str(TMP / "OUT")
cli_module.RESOURCES_FILE = trust
cli_module.DOC_NAME = "Routing Test"
cli_module.LAYER_LABEL = "Fictional Trust (2001)"

cli_module.AMENDMENTS = []
cli_module.main()
check(calls == ["single"], "no amendments: ONLY the single-document call ran")
single_out = TMP / "OUT" / "Routing_Test_annotation.xlsx"
check(single_out.exists(), "single pipeline wrote its workbook")
check("Amendment 1" not in str(load_workbook(single_out)["Annotation"]["G11"].value or ""),
      "single pipeline used the v0 template (no amendment columns)")

calls.clear()
cli_module.AMENDMENTS = [
    {"number": 1, "label": "First Amendment (2003)", "path": a1},
    {"number": 2, "label": "Second Amendment (2004)", "path": a2},
    {"number": 3, "label": "Third Amendment (2006)", "path": str(TMP / "a3.txt")},
]
cli_module.main()
check(calls == ["multi"], "any amendments: ONLY the packet call ran")
multi_out = TMP / "OUT" / "Routing_Test_annotation_with_amendments.xlsx"
check(multi_out.exists(), "multi pipeline wrote its workbook")
check(load_workbook(multi_out)["Annotation"]["H11"].value.startswith("Second Amendment"),
      "multi pipeline used the amendments template")

# Both renderers write to memory as well as to a path (needed where there is no disk).
import io
from render import render_to_workbook
buf = io.BytesIO()
render_to_workbook(tiny_single_annotation(), ANNOTATION_TEMPLATE_FILE, buf)
check(len(buf.getvalue()) > 5000 and load_workbook(io.BytesIO(buf.getvalue()))["Annotation"]["B4"].value == "stub",
      "single renderer writes a valid workbook into a BytesIO")
buf = io.BytesIO()
render_packet_to_workbook(make_packet_annotation(), pk, AMENDMENT_TEMPLATE_FILE, buf)
check(load_workbook(io.BytesIO(buf.getvalue()))["Annotation"]["H11"].value.startswith("Second Amendment"),
      "multi renderer writes a valid workbook into a BytesIO")

# The single-document renderer now writes the dropdown's exact 'Ambiguous' string.
amb = tiny_single_annotation()
amb.bequest_language.trigger_event = SingleFinding(status="Ambiguous", answer="a / b", citation="p. 1", note="n")
buf = io.BytesIO()
render_to_workbook(amb, ANNOTATION_TEMPLATE_FILE, buf)
check(load_workbook(io.BytesIO(buf.getvalue()))["Annotation"]["D19"].value == "Ambiguous \u2014 see note",
      "single renderer maps Ambiguous to the template's dropdown string")

# ---------------------------------------------------------------------------
# 7. The packet request builds, and the packet text reads in order
# ---------------------------------------------------------------------------
from llm_multi import build_packet_request, render_packet
req = build_packet_request(pk, model="claude-sonnet-4-6", effort="high", max_tokens=48000,
                           rules_path=SETTINGS.rules_multi)
check(req["tools"][0]["strict"] is True and req["tool_choice"]["name"] == "analyze_documents",
      "packet request: strict tool, forced tool choice")
text = render_packet(pk)
check(text.index("ORIGINAL TRUST") < text.index("AMENDMENT 1") < text.index("AMENDMENT 2") < text.index("AMENDMENT 3"),
      "packet text is in reading order")
gap = build_packet("T", "Trust", trust, [{"number": 1, "label": "F", "path": a1}, {"number": 3, "label": "T", "path": a2}])
check("NOT PROVIDED: amendment(s) [2]" in render_packet(gap), "a gap in the chain is stated to the model")

print(f"\nALL {passed} CHECKS PASSED")
