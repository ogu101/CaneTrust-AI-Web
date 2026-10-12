"""
Renders a PacketAnnotation into a copy of MiLA_Annotation_Template_Amendments_v0.xlsx.

Plain, deterministic Python, like render.py. What is different here:
  * three column blocks per row (original / one per amendment / operative);
  * unused amendment columns are hidden and the used ones are headed with the
    real amendment label, its form, and whether a later amendment revoked it;
  * sections that outgrow the template's placeholder rows are extended
    correctly: every later section is shifted by the rows already inserted,
    and the row heights, the Status dropdowns and the per-amendment tally
    formulas are moved/extended to match;
  * row heights are grown to fit long answers (the template's fixed heights
    would otherwise clip verbatim quotations).
"""

import copy
import datetime as dt
import math
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange

from cell_map_multi import (
    HEADER,
    FIXED_ROWS,
    VARIABLE_SECTIONS,
    ORIGINAL_COLS,
    OPERATIVE_COLS,
    AMENDMENT_FIRST_COL,
    AMENDMENT_SLOTS,
    AMENDMENT_HEADER_ROW,
    LAST_COL,
    STATUS_TO_CELL,
    FORM_TO_TEXT,
    FIXED_INPUT_RANGES,
)
from packet import Packet
from workbook_io import save_workbook
from schema_multi import PacketAnnotation, FieldHistory, AmendmentEntry

BAND_FILL = "FFEDF2F0"      # section header rows; autofit skips them
LINE_HEIGHT = 12.75         # points per line of Arial 10
MAX_ROW_HEIGHT = 409.5      # Excel's limit


# ----------------------------------------------------------------------------
# small helpers
# ----------------------------------------------------------------------------

def _amendment_col(index: int) -> str:
    return get_column_letter(AMENDMENT_FIRST_COL + index)


def _amendment_text(entry: AmendmentEntry) -> str:
    if entry.effect == "No change":
        return "No change"
    text = f"{entry.effect}:"
    if entry.description.strip():
        text += f" {entry.description.strip()}"
    if entry.citation.strip():
        text += f" ({entry.citation.strip()})"
    return text


def _write_block(ws, row: int, finding, cols: dict) -> None:
    ws[f"{cols['answer']}{row}"] = finding.answer
    ws[f"{cols['citation']}{row}"] = finding.citation
    ws[f"{cols['status']}{row}"] = STATUS_TO_CELL.get(finding.status, finding.status)
    ws[f"{cols['note']}{row}"] = finding.note


def _write_history(ws, row: int, history: FieldHistory, amendment_numbers: list[int]) -> None:
    _write_block(ws, row, history.original, ORIGINAL_COLS)
    by_number = {e.amendment_number: e for e in history.amendments}
    for index, number in enumerate(amendment_numbers):
        entry = by_number.get(number)
        ws[f"{_amendment_col(index)}{row}"] = (
            _amendment_text(entry) if entry else
            "Could not determine: no entry for this amendment."
        )
    _write_block(ws, row, history.operative, OPERATIVE_COLS)


def _copy_row_style(ws, src_row: int, dst_row: int) -> None:
    for col in range(1, LAST_COL + 1):
        src, dst = ws.cell(row=src_row, column=col), ws.cell(row=dst_row, column=col)
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format


def _insert_rows(ws, at_row: int, amount: int, style_src_row: int) -> None:
    """ws.insert_rows moves cells but not row heights, so carry those over, and
    give the new rows the look of the row above."""
    heights = {r: ws.row_dimensions[r].height for r in range(at_row, ws.max_row + 1)}
    ws.insert_rows(at_row, amount)
    for r, h in heights.items():
        ws.row_dimensions[r + amount].height = h
    for i in range(amount):
        ws.row_dimensions[at_row + i].height = ws.row_dimensions[style_src_row].height
        _copy_row_style(ws, style_src_row, at_row + i)


def _ensure_capacity(ws, key: str, n_items: int, offset: int):
    """Make room for n_items in a variable section that starts `offset` rows
    below where the original template has it. Returns (first_row, total_rows,
    rows_inserted)."""
    cfg = VARIABLE_SECTIONS[key]
    first = cfg["first_row"] + offset
    capacity = cfg["placeholder_count"]
    inserted = max(0, n_items - capacity)
    if inserted:
        _insert_rows(ws, first + capacity, inserted, style_src_row=first + capacity - 1)
    return first, capacity + inserted, inserted


def _ungroup_columns(ws) -> None:
    """Spreadsheet apps save runs of equal-width columns as ONE <col min=7 max=14>
    entry, which openpyxl loads under the first column's key only. Touching a
    single column inside such a run (to hide it) would write overlapping column
    definitions and make Excel offer to 'repair' the file. Split every run into
    one entry per column first."""
    for dim in list(ws.column_dimensions.values()):
        if dim.min and dim.max and dim.max > dim.min:
            low, high = dim.min, dim.max
            dim.max = low
            for index in range(low + 1, high + 1):
                letter = get_column_letter(index)
                ws.column_dimensions[letter].width = dim.width
                ws.column_dimensions[letter].hidden = dim.hidden


def _autofit_row(ws, row: int, text_cols: list[str]) -> None:
    if ws[f"A{row}"].fill.fgColor.rgb == BAND_FILL:      # section header rows
        return
    lines = 1
    for col in text_cols:
        value = ws[f"{col}{row}"].value
        if not value:
            continue
        width = ws.column_dimensions[col].width or 8.43
        per_line = max(1, int(width))
        needed = sum(max(1, math.ceil(len(part) / per_line)) for part in str(value).split("\n"))
        lines = max(lines, needed)
    current = ws.row_dimensions[row].height or 15
    ws.row_dimensions[row].height = min(MAX_ROW_HEIGHT, max(current, lines * LINE_HEIGHT + 4))


# ----------------------------------------------------------------------------
# main entry point
# ----------------------------------------------------------------------------

def render_packet_to_workbook(
    annotation: PacketAnnotation,
    packet: Packet,
    template_path: str,
    out,
    annotator: str = "MiLA Assistant (LLM draft — attorney review required)",
) -> None:
    amendment_numbers = packet.amendment_numbers
    n_amendments = len(amendment_numbers)
    if not 1 <= n_amendments <= AMENDMENT_SLOTS:
        raise ValueError(f"Expected 1-{AMENDMENT_SLOTS} amendments, got {n_amendments}")

    wb = load_workbook(template_path)
    ws = wb["Annotation"]
    _ungroup_columns(ws)

    # ---- header block --------------------------------------------------------
    ws[HEADER["document_name"]] = annotation.document_name
    ws[HEADER["annotator"]] = annotator
    ws[HEADER["date_completed"]] = dt.date.today().isoformat()

    # ---- amendment column headers; hide the slots this packet doesn't use ----
    info_by_number = {i.amendment_number: i for i in annotation.instruments}
    revoked_by: dict[int, int] = {}
    for info in annotation.instruments:
        for revoked in info.revokes:
            revoked_by[revoked] = info.amendment_number

    for index, amendment in enumerate(packet.amendments):
        info = info_by_number.get(amendment.number)
        form = FORM_TO_TEXT.get(info.form, info.form) if info else "form unclear"
        header = f"{amendment.label}\n{form}"
        if amendment.number in revoked_by:
            header += f" — REVOKED by Amendment {revoked_by[amendment.number]}"
        ws.cell(row=AMENDMENT_HEADER_ROW, column=AMENDMENT_FIRST_COL + index).value = header
    ws.row_dimensions[AMENDMENT_HEADER_ROW].height = 48
    for index in range(n_amendments, AMENDMENT_SLOTS):
        ws.column_dimensions[_amendment_col(index)].hidden = True

    # ---- sections 1-3: one row per field -------------------------------------
    for path, row in FIXED_ROWS.items():
        obj = annotation
        for part in path.split("."):
            obj = getattr(obj, part)
        _write_history(ws, row, obj, amendment_numbers)

    # ---- sections 4-8: variable length, each shifted by rows inserted above --
    offset = 0

    first, total, inserted = _ensure_capacity(ws, "defined_terms", len(annotation.defined_terms), offset)
    terms_first, terms_total = first, total
    for i in range(total):
        row = first + i
        if i < len(annotation.defined_terms):
            term = annotation.defined_terms[i]
            ws[f"A{row}"] = term.term
            _write_history(ws, row, term.history, amendment_numbers)
        else:
            ws[f"{ORIGINAL_COLS['status']}{row}"] = STATUS_TO_CELL["Not in this document"]
            ws[f"{OPERATIVE_COLS['status']}{row}"] = STATUS_TO_CELL["Not in this document"]
    offset += inserted

    first, total, inserted = _ensure_capacity(ws, "um_mentions", len(annotation.um_mentions), offset)
    mentions_first, mentions_total = first, total
    for i in range(total):
        row = first + i
        if i < len(annotation.um_mentions):
            ws[f"A{row}"] = f"Mention {i + 1}"
            _write_history(ws, row, annotation.um_mentions[i], amendment_numbers)
        else:
            ws[f"{ORIGINAL_COLS['status']}{row}"] = STATUS_TO_CELL["Not in this document"]
            ws[f"{OPERATIVE_COLS['status']}{row}"] = STATUS_TO_CELL["Not in this document"]
    offset += inserted
    last_history_row = mentions_first + mentions_total - 1

    first, total, inserted = _ensure_capacity(
        ws, "found_but_uncategorized", len(annotation.found_but_uncategorized), offset)
    for i, item in enumerate(annotation.found_but_uncategorized):
        row = first + i
        ws[f"A{row}"] = f"{i + 1}."
        ws[f"B{row}"] = item.text
        ws[f"C{row}"] = item.citation
        ws[f"E{row}"] = item.note
    offset += inserted

    for key, items in (
        ("judgment_calls", annotation.judgment_calls),
        ("questions_for_client", annotation.questions_for_client),
    ):
        first, total, inserted = _ensure_capacity(ws, key, len(items), offset)
        for i, text in enumerate(items):
            ws[f"A{first + i}"] = f"{i + 1}."
            ws[f"B{first + i}"] = text
        offset += inserted

    # ---- keep the template's machinery pointing at the right rows ------------
    # Status dropdowns (original D and operative R columns).
    ranges = list(FIXED_INPUT_RANGES) + [
        (terms_first, terms_first + terms_total - 1),
        (mentions_first, mentions_first + mentions_total - 1),
    ]
    for dv in ws.data_validations.dataValidation:
        col = re.match(r"[A-Z]+", str(dv.sqref)).group(0)
        dv.sqref = MultiCellRange(" ".join(f"{col}{lo}:{col}{hi}" for lo, hi in ranges))

    # "Fields changed or flagged" tally above each amendment column.
    for index in range(n_amendments):
        col = _amendment_col(index)
        rng = f"{col}13:{col}{last_history_row}"
        ws[f"{col}9"] = f'=COUNTIF({rng},"*")-COUNTIF({rng},"No change*")'

    # ---- row heights ---------------------------------------------------------
    text_cols = (
        list(ORIGINAL_COLS.values())
        + [_amendment_col(i) for i in range(n_amendments)]
        + list(OPERATIVE_COLS.values())
    )
    for row in range(13, ws.max_row + 1):
        _autofit_row(ws, row, text_cols)

    save_workbook(wb, out)
