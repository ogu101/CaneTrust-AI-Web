"""
API-free tests for extraction and the web API.

    python test_web.py

The model is replaced with canned annotations, so no API key or network is needed. The PDFs
are the FICTIONAL Marlowe packet in fixtures/.
"""

import base64
import contextlib
import io
import json
import os
import sys
import tempfile
import time
from pathlib import Path

os.environ.setdefault("ANTHROPIC_API_KEY", "test-key-not-used")
for name in ("VERCEL", "CANETRUST_ACCESS_CODE", "CANETRUST_MAX_UPLOAD_MB", "CANETRUST_DEADLINE_SECONDS"):
    os.environ.pop(name, None)

import anthropic
import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

import service
from api import index as api_index
from extract import ExtractionError, extract_text
from schema import DocumentAnnotation
from schema_multi import reconcile_with_packet
from testkit import make_packet_annotation, tiny_single_annotation

FIX = Path(__file__).parent / "fixtures"
passed = 0


def check(condition, message):
    global passed
    assert condition, f"FAILED: {message}"
    passed += 1


@contextlib.contextmanager
def env(**values):
    old = {k: os.environ.get(k) for k in values}
    for k, v in values.items():
        if v is None:
            os.environ.pop(k, None)
        else:
            os.environ[k] = v
    try:
        yield
    finally:
        for k, v in old.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def pdf(name):
    return (FIX / name).read_bytes()


def blank_pdf(pages=1):
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def raises_kind(fn, kind):
    try:
        fn()
    except ExtractionError as e:
        return e.kind == kind, str(e)
    return False, "did not raise"


# ---------------------------------------------------------------------------
# 1. Extraction
# ---------------------------------------------------------------------------
TRUST = pdf("00_marlowe_trust_original.pdf")

ex = extract_text("trust.pdf", TRUST)
check(ex.pages == 2 and ex.method == "pdf-text-layer" and not ex.warnings, "a text PDF extracts cleanly")
check(ex.text.startswith("[PDF page 1 of 2]") and "[PDF page 2 of 2]" in ex.text, "page markers are inserted")
check(ex.text.index("[PDF page 2 of 2]") > ex.text.index("4.3 Distribution of the Residue"),
      "markers fall at the real page boundary (4.3 starts on page 1)")
check("UNIVERSITY OF MIAMI" in ex.text, "the document's own text survives")
check("[PDF page" not in extract_text("t.pdf", TRUST, page_markers=False).text, "markers can be turned off")

ok, msg = raises_kind(lambda: extract_text("scan.pdf", blank_pdf(3)), "no_text_layer")
check(ok and "scan" in msg and "scan.pdf" in msg, "a PDF with no text layer is refused, by name, as a scan")

writer = PdfWriter()
for page in PdfReader(io.BytesIO(TRUST)).pages:
    writer.add_page(page)
writer.add_blank_page(width=612, height=792)
mixed = io.BytesIO()
writer.write(mixed)
ex = extract_text("mixed.pdf", mixed.getvalue())
check("[PDF page 3 of 3: no extractable text]" in ex.text, "an unreadable page is marked in the text itself")
check(len(ex.warnings) == 1 and "page(s) 3" in ex.warnings[0], "and reported as a warning")

writer = PdfWriter()
for page in PdfReader(io.BytesIO(TRUST)).pages:
    writer.add_page(page)
writer.encrypt("secret")
locked = io.BytesIO()
writer.write(locked)
check(raises_kind(lambda: extract_text("locked.pdf", locked.getvalue()), "encrypted")[0], "a password-protected PDF is refused")
check(raises_kind(lambda: extract_text("fake.pdf", b"hello"), "invalid")[0], "a non-PDF named .pdf is refused")
check(raises_kind(lambda: extract_text("brief.docx", b"PK..."), "unsupported")[0], "an unsupported type is refused")
check(raises_kind(lambda: extract_text("e.txt", b"  \n"), "empty")[0], "an empty .txt is refused")
check(raises_kind(lambda: extract_text("t.pdf", TRUST, max_pages=1), "too_long")[0], "a PDF over the page limit is refused")
check(extract_text("q.txt", "\u201cquoted\u201d".encode("cp1252")).text == "\u201cquoted\u201d", "cp1252 text decodes")
check(extract_text("u.txt", "caf\u00e9".encode("utf-8")).text == "caf\u00e9", "utf-8 text decodes")

# ---------------------------------------------------------------------------
# 2. The API, with the model stubbed
# ---------------------------------------------------------------------------
client = TestClient(api_index.app)
calls: list = []
SLOW = {"seconds": 0.0}


STATUS_SEEN: dict = {}


def fake_single(name, layers, **kw):
    STATUS_SEEN["single"] = layers[0].status
    calls.append(("single", name, [layer.label for layer in layers], layers[0].text))
    time.sleep(SLOW["seconds"])
    return tiny_single_annotation()


def fake_multi(packet, **kw):
    calls.append(("multi", packet.name, packet.amendment_numbers, packet.trust.text))
    time.sleep(SLOW["seconds"])
    annotation = make_packet_annotation(n_terms=2, n_mentions=2)
    reconcile_with_packet(annotation, packet.amendment_numbers)      # as the real analyze_packet does
    return annotation


service.analyze_document = fake_single
service.analyze_packet = fake_multi

AMENDMENT_FILES = [
    ("01_marlowe_first_amendment.pdf", 1, "First Amendment (2004)"),
    ("02_marlowe_second_amendment.pdf", 2, "Second Amendment (2009)"),
    ("03_marlowe_third_amendment_and_restatement.pdf", 3, "Third Amendment and Restatement (2013)"),
    ("04_marlowe_fourth_amendment.pdf", 4, "Fourth Amendment (2016)"),
]


def post(trust=("00_marlowe_trust_original.pdf", TRUST), amendments=(), data=None, headers=None, meta=None):
    files = [("trust", (trust[0], trust[1], "application/pdf"))]
    for fname, content, _n, _label in amendments:
        files.append(("amendment_files", (fname, content, "application/pdf")))
    if meta is None:
        meta = [{"number": n, "label": label} for _f, _c, n, label in amendments]
    form = {"amendments_meta": json.dumps(meta), **(data or {})}
    return client.post("/api/annotate", files=files, data=form,
                       headers={"Accept": "application/x-ndjson", **(headers or {})})


def events(response):
    return [json.loads(line) for line in response.text.splitlines() if line.strip()]


def amendment(i):
    fname, number, label = AMENDMENT_FILES[i]
    return (fname, pdf(fname), number, label)


# -- health
h = client.get("/api/health")
check(h.status_code == 200 and h.json()["ok"] and h.json()["assets_present"], "health: assets are bundled")
check("test-key-not-used" not in h.text and h.json()["api_key_set"] is True, "health reports the key is set, never the key")

# -- multi: all four amendments, uploaded OUT of order
calls.clear()
r = post(amendments=[amendment(3), amendment(1), amendment(0), amendment(2)])
ev = events(r)
check(r.status_code == 200 and r.headers["content-type"].startswith("application/x-ndjson"), "success streams NDJSON")
check(ev[0] == {"type": "progress", "message": "Documents accepted."}, "first event: documents accepted")
check(any(e["type"] == "progress" and "4 amendment(s)" in e["message"] and "multi" in e["message"] for e in ev),
      "progress reports the triage decision")
done = ev[-1]
check(done["type"] == "done" and done["route"] == "multi", "last event is 'done', routed multi")
check(done["filename"] == "00_marlowe_trust_original_annotation_with_amendments.xlsx", "output filename is derived from the trust")
check(calls and calls[0][0] == "multi" and calls[0][2] == [1, 2, 3, 4], "amendments reach the model sorted by number, not upload order")
check("[PDF page 1 of 2]" in calls[0][3] and "UNIVERSITY OF MIAMI" in calls[0][3], "the model receives extracted, page-marked text")

wb = load_workbook(io.BytesIO(base64.b64decode(done["xlsx_base64"])))
ws = wb["Annotation"]
check(ws["G11"].value.startswith("First Amendment (2004)") and ws["J11"].value.startswith("Fourth Amendment (2016)"),
      "the workbook's amendment columns run 1..4 with their labels")
check(ws.column_dimensions["K"].hidden and ws.column_dimensions["N"].hidden and not ws.column_dimensions["J"].hidden,
      "unused amendment columns are hidden")

# -- single: trust only
calls.clear()
r = post()
ev = events(r)
check(ev[-1]["type"] == "done" and ev[-1]["route"] == "single", "no amendments: routed single")
check(calls == [("single", "00_marlowe_trust_original", ["00_marlowe_trust_original"], calls[0][3])],
      "single path: the document name and label default to the file name")
wb = load_workbook(io.BytesIO(base64.b64decode(ev[-1]["xlsx_base64"])))
check(wb["Annotation"]["B4"].value == "stub" and wb["Annotation"]["G11"].value in (None, ""),
      "single path used the v0 template")
check(ev[-1]["filename"] == "00_marlowe_trust_original_annotation.xlsx", "single output filename")

# -- explicit names are honoured and the filename is made safe
r = post(data={"doc_name": "Marlowe / Trust: 1998?", "trust_label": "Marlowe Trust (1998)"})
check(events(r)[-1]["filename"] == "Marlowe__Trust_1998_annotation.xlsx", "unsafe characters are stripped from the download name")
check(calls[-1][1] == "Marlowe / Trust: 1998?" and calls[-1][2] == ["Marlowe Trust (1998)"], "doc_name and trust_label reach the pipeline")

# -- plain .txt uploads work too
r = post(trust=("trust.txt", b"The trust text. UNIVERSITY OF MIAMI."))
check(events(r)[-1]["type"] == "done", ".txt uploads are accepted")

# -- a gap in the chain is reported in progress
r = post(amendments=[amendment(0), amendment(3)])
check(any("[2, 3]" in e.get("message", "") and "not supplied" in e["message"] for e in events(r)),
      "a gap in the amendment numbers is announced")

# -- access control
with env(CANETRUST_ACCESS_CODE="sesame"):
    r = post()
    check(r.status_code == 401 and r.json()["code"] == "bad_access_code", "no access code -> 401")
    check(post(data={"access_code": "wrong"}).status_code == 401, "wrong access code -> 401")
    check(events(post(data={"access_code": "sesame"}))[-1]["type"] == "done", "right code in the form -> allowed")
    check(events(post(headers={"X-Access-Code": "sesame"}))[-1]["type"] == "done", "right code in a header -> allowed")
with env(VERCEL="1"):
    r = post()
    check(r.status_code == 503 and r.json()["code"] == "not_configured", "deployed with no access code configured: refuses to run open")
    check(client.get("/api/health").json()["access_code_required"] is False, "health shows the gate is not configured")

# -- validation errors: proper HTTP statuses, before anything streams
r = post(trust=("scan.pdf", blank_pdf(2)))
check(r.status_code == 422 and r.json()["code"] == "no_text_layer" and "scan.pdf" in r.json()["error"], "a scanned trust -> 422, naming the file")
r = post(amendments=[amendment(0)], meta=[{"number": 1, "label": "A"}, {"number": 2, "label": "B"}])
check(r.status_code == 400 and r.json()["code"] == "bad_request", "metadata count != file count -> 400")
r = post(amendments=[amendment(0), amendment(1)], meta=[{"number": 1, "label": "A"}, {"number": 1, "label": "B"}])
check(r.status_code == 400 and "more than once" in r.json()["error"], "duplicate amendment numbers -> 400")
r = post(amendments=[amendment(0)], meta=[{"number": "x", "label": "A"}])
check(r.status_code == 400 and r.json()["code"] == "bad_amendment_number", "a non-numeric amendment number -> 400")
r = post(amendments=[amendment(0)], meta=[{"number": "2", "label": ""}])
check(events(r)[-1]["type"] == "done" and calls[-1][2] == [2], "a numeric string is accepted as the number")
r = post(trust=("notes.docx", b"PK\x03\x04"))
check(r.status_code == 415 and r.json()["code"] == "unsupported", "an unsupported file type -> 415")
r = client.post("/api/annotate", data={"amendments_meta": "not json"}, files=[("trust", ("t.txt", b"x", "text/plain"))])
check(r.status_code == 400, "malformed metadata -> 400")
r = client.post("/api/annotate", data={})
check(r.status_code == 422, "no trust file at all -> rejected")
many = [(f"a{i}.txt", b"Amendment text.", i, f"A{i}") for i in range(1, 10)]
r = post(amendments=many)
check(r.status_code == 400 and r.json()["code"] == "too_many_amendments", "more than 8 amendments -> 400")
with env(CANETRUST_MAX_UPLOAD_MB="0.001"):
    r = post()
    check(r.status_code == 413 and r.json()["code"] == "too_large", "an oversized upload -> 413")
with env(CANETRUST_MAX_CHARS="500"):
    r = post()
    check(r.status_code == 400 and "character limit" in r.json()["error"], "too much text for one packet -> 400")

# -- model failures arrive as a final 'error' event, with safe messages
def failing(exc):
    def _raise(*a, **k):
        raise exc
    return _raise


request = httpx.Request("POST", "https://api.anthropic.com/v1/messages")
SECRET = "boom-SECRET-document-text"
cases = [
    (anthropic.RateLimitError("slow down", response=httpx.Response(429, request=request), body=None), "rate_limited"),
    (anthropic.AuthenticationError("bad key sk-ant-xyz", response=httpx.Response(401, request=request), body=None), "model_auth"),
    (anthropic.APIConnectionError(request=request), "model_unreachable"),
    (anthropic.InternalServerError("oops", response=httpx.Response(500, request=request), body=None), "model_error"),
    (RuntimeError("Output was cut off at max_tokens=8000"), "run_failed"),
    (ValueError(SECRET), "internal"),
]
try:
    DocumentAnnotation.model_validate({"document_name": SECRET})
except Exception as validation_error:
    cases.append((validation_error, "bad_model_output"))

stderr = io.StringIO()
for exc, code in cases:
    service.analyze_document = failing(exc)
    with contextlib.redirect_stderr(stderr):
        r = post()
    ev = events(r)
    check(r.status_code == 200 and ev[-1]["type"] == "error" and ev[-1]["code"] == code,
          f"model failure {type(exc).__name__} -> error event {code!r}")
    shown = ev[-1]["message"]
    check(SECRET not in shown and "sk-ant" not in shown, f"{code}: the user-facing message leaks nothing")
    if code == "run_failed":
        check("Output was cut off" in shown, "a crafted RuntimeError message is passed through to the user")
logged = stderr.getvalue()
check("[error] ValueError: (message withheld)" in logged and SECRET not in logged and "sk-ant" not in logged,
      "server logs name the error class but never the message of an unexpected exception")
check("[error] RuntimeError: Output was cut off" in logged, "messages we wrote ourselves are logged")

service.analyze_document = fake_single
with env(ANTHROPIC_API_KEY=None):
    ev = events(post())
    check(ev[-1]["type"] == "error" and "ANTHROPIC_API_KEY" in ev[-1]["message"], "a missing API key is reported plainly")

# -- long runs: heartbeats keep the stream alive, and a deadline ends it cleanly
api_index.HEARTBEAT_SECONDS = 0.15
SLOW["seconds"] = 0.7
ev = events(post())
beats = [e for e in ev if e["type"] == "heartbeat"]
check(len(beats) >= 2 and ev[-1]["type"] == "done", "a slow run emits heartbeats, then completes")
with env(CANETRUST_DEADLINE_SECONDS="1"):
    SLOW["seconds"] = 3.0
    ev = events(post())
    check(ev[-1]["type"] == "error" and ev[-1]["code"] == "deadline", "past the deadline the stream ends with a clear error")
SLOW["seconds"] = 0.0
api_index.HEARTBEAT_SECONDS = 10

# ---------------------------------------------------------------------------
# 2b. Classic mode: the original contract (file/label/status in, workbook bytes out)
# ---------------------------------------------------------------------------
service.analyze_document = fake_single
service.analyze_packet = fake_multi


def classic(file=("t.pdf", TRUST), data=None, headers=None):
    return client.post("/api/annotate", files=[("file", (file[0], file[1], "application/pdf"))],
                       data=data or {}, headers=headers or {})


calls.clear()
r = classic(data={"label": "Trust document", "status": "OPERATIVE"})
check(r.status_code == 200 and r.headers["content-type"] == api_index.XLSX, "classic: the workbook itself comes back")
check(r.headers["content-disposition"].startswith('attachment; filename="t_annotation.xlsx"'), "classic: download name in Content-Disposition")
check(r.headers["x-canetrust-route"] == "single", "classic: the route is reported in a header")
check(load_workbook(io.BytesIO(r.content))["Annotation"]["B4"].value == "stub", "classic: the bytes are a real workbook")
check(calls[-1][2] == ["Trust document"] and STATUS_SEEN["single"] == "OPERATIVE", "classic: label and status reach the pipeline")
classic(data={"status": "IN FORCE"})
check(STATUS_SEEN["single"] == "IN FORCE", "classic: a custom status is honoured")
classic()
check(STATUS_SEEN["single"] == "OPERATIVE", "classic: status defaults to OPERATIVE")

r = classic(file=("Tr\u00fcst \u2013 1998.pdf", TRUST))
check(r.status_code == 200 and "filename*=UTF-8''" in r.headers["content-disposition"]
      and r.headers["content-disposition"].isascii(), "classic: a non-ASCII file name is safe in the header")

r = client.post("/api/annotate", data={"label": "x"})
check(r.status_code == 422 and r.json()["code"] == "missing_file", "classic: no file -> 422 as JSON")
r = classic(file=("scan.pdf", blank_pdf(1)))
check(r.status_code == 422 and r.json()["code"] == "no_text_layer", "classic: a scan -> 422 with the explanation")

classic_failures = [
    (anthropic.RateLimitError("slow", response=httpx.Response(429, request=httpx.Request("POST", "https://x")), body=None), 429, "rate_limited"),
    (anthropic.AuthenticationError("bad key sk-ant-q", response=httpx.Response(401, request=httpx.Request("POST", "https://x")), body=None), 502, "model_auth"),
    (RuntimeError("Output was cut off at max_tokens=8000"), 500, "run_failed"),
    (ValueError("boom-SECRET-document-text"), 500, "internal"),
]
sink = io.StringIO()
for exc, status_code, code in classic_failures:
    def _boom(*a, _e=exc, **k):
        raise _e
    service.analyze_document = _boom
    with contextlib.redirect_stderr(sink):
        r = classic()
    check(r.status_code == status_code and r.json()["code"] == code, f"classic: {type(exc).__name__} -> HTTP {status_code}")
    check("SECRET" not in r.text and "sk-ant" not in r.text, f"classic {code}: nothing sensitive in the response")
service.analyze_document = fake_single

with env(CANETRUST_ACCESS_CODE="sesame"):
    check(classic().status_code == 401, "classic: the access code is enforced here too")
    check(classic(headers={"X-Access-Code": "sesame"}).status_code == 200, "classic: the code is accepted in a header")
with env(VERCEL="1"):
    check(classic().status_code == 503, "classic: a deployment with no code refuses to run open")
with env(VERCEL="1", CANETRUST_ALLOW_OPEN="1"):
    check(classic().status_code == 200, "CANETRUST_ALLOW_OPEN lets a code-less deployment run (explicit opt-in)")
    check(client.get("/api/health").json()["open_access_allowed"] is True, "health reports that open access is on")
check(client.get("/api/health").json()["open_access_allowed"] is False, "open access is off by default")

# ---------------------------------------------------------------------------
# 3. The command-line tool extracts PDFs itself (no separate extraction step)
# ---------------------------------------------------------------------------
import cli as cli_module

seen = {}
service.analyze_document = lambda name, layers, **kw: seen.update(single_text=layers[0].text) or tiny_single_annotation()
service.analyze_packet = lambda packet, **kw: seen.update(multi=(packet.trust.text, [a.text for a in packet.amendments])) or (
    lambda a: (reconcile_with_packet(a, packet.amendment_numbers), a)[1])(make_packet_annotation(n_terms=1, n_mentions=1))

out_dir = tempfile.mkdtemp()
cli_module.OUTPUT_DIR = out_dir
cli_module.DOC_NAME = "CLI PDF Test"
cli_module.LAYER_LABEL = "Marlowe Trust (1998)"
cli_module.RESOURCES_FILE = str(FIX / "00_marlowe_trust_original.pdf")

cli_module.AMENDMENTS = []
with contextlib.redirect_stdout(io.StringIO()):
    cli_module.main()
check("[PDF page 1 of 2]" in seen["single_text"] and "UNIVERSITY OF MIAMI" in seen["single_text"],
      "cli: a PDF named in config.py is extracted by the pipeline, with page markers")
check((Path(out_dir) / "CLI_PDF_Test_annotation.xlsx").exists(), "cli: single-path workbook written")

cli_module.AMENDMENTS = [
    {"number": 2, "label": "Second", "path": str(FIX / "02_marlowe_second_amendment.pdf")},
    {"number": 1, "label": "First", "path": str(FIX / "01_marlowe_first_amendment.pdf")},
]
with contextlib.redirect_stdout(io.StringIO()):
    cli_module.main()
trust_text, amendment_texts = seen["multi"]
check("Declaration of Trust dated March 12, 1998" in trust_text and "[PDF page 1 of 1]" in amendment_texts[0],
      "cli: amendment PDFs are extracted too")
check("FIRST AMENDMENT" in amendment_texts[0].upper() and "SECOND AMENDMENT" in amendment_texts[1].upper(),
      "cli: amendments reach the model sorted by number")
check((Path(out_dir) / "CLI_PDF_Test_annotation_with_amendments.xlsx").exists(), "cli: multi-path workbook written")

cli_module.RESOURCES_FILE = str(FIX / "scan_like_blank.pdf")
try:
    cli_module.main()
    check(False, "cli: a scanned PDF should stop the run")
except ValueError as exc:
    check("no text layer" in str(exc) and "Original trust" in str(exc), "cli: a scanned PDF stops the run with the same explanation")

# -- nothing is written to the working directory by a run
before = set(Path(".").rglob("*.xlsx"))
events(post())
check(set(Path(".").rglob("*.xlsx")) == before, "a run writes no files (the result travels in memory)")

print(f"\nALL {passed} CHECKS PASSED")
